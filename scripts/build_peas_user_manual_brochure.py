from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
DOCX_PATH = OUT / "docx" / "PeAS_User_Manual_Brochure.docx"
SCREEN = OUT / "evidence" / "user_manual_screenshots"
CROPS = OUT / "evidence" / "user_manual_screenshots" / "brochure_crops"
ASSETS = ROOT / "Deno" / "Public" / "Components" / "images"

GREEN = RGBColor(0x00, 0x6B, 0x4E)
DEEP_GREEN = RGBColor(0x00, 0x46, 0x36)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
INK = RGBColor(0x1F, 0x2B, 0x2A)
MUTED = RGBColor(0x5E, 0x6B, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE_GREEN = "EAF3EE"
PALE_GOLD = "FBF4D8"
PALE_BLUE = "EEF4F7"
LINE = "D4DFD9"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size="4"):
    borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_no_borders(table):
    borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def set_cell_width(cell, width_inches: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    """Set an explicit DXA table width and matching grid for predictable Word layout."""
    table.autofit = False
    total = sum(int(width * 1440) for width in widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    grid_cols = list(grid)
    while len(grid_cols) < len(widths):
        col = OxmlElement("w:gridCol")
        grid.append(col)
        grid_cols.append(col)
    for col, width in zip(grid_cols, widths):
        col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)


def set_run(run, size=9.4, color=INK, bold=False, italic=False, font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def set_para(p, before=0, after=4, line=1.06, keep=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.widow_control = True
    p.paragraph_format.keep_with_next = keep


def add_text(doc_or_cell, text, size=9.4, color=INK, bold=False, italic=False, align=None, before=0, after=4, style=None):
    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    set_para(p, before, after)
    r = p.add_run(text)
    set_run(r, size, color, bold, italic)
    return p


def add_rich(doc_or_cell, parts, size=9.4, align=None, before=0, after=4):
    p = doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    set_para(p, before, after)
    for text, kwargs in parts:
        r = p.add_run(text)
        set_run(r, size=size, **kwargs)
    return p


def add_label(doc_or_cell, text, color=GOLD, size=8.2, after=2):
    p = doc_or_cell.add_paragraph()
    set_para(p, after=after, keep=True)
    r = p.add_run(text.upper())
    set_run(r, size=size, color=color, bold=True)
    return p


def add_heading(doc_or_cell, text, level=1):
    p = doc_or_cell.add_paragraph()
    set_para(p, before=0 if level == 1 else 5, after=4 if level == 1 else 2, line=1.0, keep=True)
    r = p.add_run(text)
    set_run(r, size=16 if level == 1 else 11.3, color=DEEP_GREEN if level == 1 else GREEN, bold=True)
    return p


def add_bullets(doc_or_cell, items, size=9.1, after=2):
    for item in items:
        p = doc_or_cell.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        set_para(p, after=after, line=1.02)
        r = p.add_run("• ")
        set_run(r, size, GREEN, True)
        r = p.add_run(item)
        set_run(r, size, INK)


def add_steps(doc_or_cell, steps, size=8.9):
    for idx, item in enumerate(steps, 1):
        p = doc_or_cell.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.21)
        p.paragraph_format.first_line_indent = Inches(-0.21)
        set_para(p, after=2, line=1.02)
        r = p.add_run(f"{idx}. ")
        set_run(r, size, GREEN, True)
        r = p.add_run(item)
        set_run(r, size, INK)


def add_callout(doc_or_cell, label, text, fill=PALE_GREEN, size=8.7):
    table = doc_or_cell.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=LINE, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=100, start=130, bottom=90, end=130)
    p = cell.paragraphs[0]
    set_para(p, after=2, line=1.02)
    r = p.add_run(label.upper() + "  ")
    set_run(r, 8.1, GOLD, True)
    r = p.add_run(text)
    set_run(r, size, INK)
    spacer = doc_or_cell.add_paragraph()
    set_para(spacer, after=0)
    return table


def add_image(doc_or_cell, path: Path, width: float, alt: str, caption: str | None = None, crop=None):
    if crop:
        path = crop_image(path, crop)
    p = doc_or_cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=2, line=1.0)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt)
    doc_pr.set("title", "Fictional PeAS demonstration screengrab")
    if caption:
        cp = doc_or_cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(cp, after=4, line=1.0)
        cr = cp.add_run(caption)
        set_run(cr, 7.6, MUTED, italic=True)
    return inline


def crop_image(path: Path, crop):
    CROPS.mkdir(parents=True, exist_ok=True)
    name = f"crop_{path.stem}_{crop[0]}_{crop[1]}_{crop[2]}_{crop[3]}.png"
    out = CROPS / name
    if not out.exists():
        with Image.open(path) as image:
            image.crop(crop).save(out)
    return out


def add_two_col(doc, widths=(2.45, 2.45), gap_fill="FFFFFF"):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_no_borders(table)
    set_table_geometry(table, widths)
    for cell, width in zip(table.rows[0].cells, widths):
        set_cell_width(cell, width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, top=0, start=45, bottom=0, end=45)
        set_cell_shading(cell, gap_fill)
    return table.rows[0].cells


def add_rule(doc, color=LINE):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color=color, size="5")
    cell = table.cell(0, 0)
    cell.height = Inches(0.01)
    set_cell_shading(cell, color)
    set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    return table


def configure(doc):
    sec = doc.sections[0]
    sec.page_width = Mm(148)
    sec.page_height = Mm(210)
    sec.top_margin = Mm(11)
    sec.bottom_margin = Mm(11)
    sec.left_margin = Mm(12)
    sec.right_margin = Mm(12)
    sec.header_distance = Mm(5)
    sec.footer_distance = Mm(5)
    styles = doc.styles
    normal = styles["Normal"]
    set_run(normal, 9.4, INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.06
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Caption", "Intense Quote", "List Bullet", "List Number"):
        try:
            st = styles[name]
            st.font.name = "Arial"
            st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        except KeyError:
            pass
    # Keep a portable metadata profile; the final scrub runs separately.
    props = doc.core_properties
    props.title = "PeAS User Manual - Brochure Booklet"
    props.subject = "Fictional demonstration-data user manual for the Paulinian electronic Archiving System"
    props.author = "Office of Research & Publications"
    props.keywords = "PeAS, user manual, brochure, demonstration data"
    props.comments = "All records and identities shown in screenshots are fictional fixtures."

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para(header, after=0)
    r = header.add_run("PeAS  /  USER MANUAL")
    set_run(r, 7.5, MUTED, True)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(footer, after=0)
    r = footer.add_run("Office of Research & Publications  •  Version 1.0  •  ")
    set_run(r, 7.2, MUTED)
    add_page_field(footer)


def start_page(doc, kicker, title, intro=None):
    add_label(doc, kicker)
    add_heading(doc, title, 1)
    if intro:
        p = doc.add_paragraph()
        set_para(p, after=6, line=1.06)
        r = p.add_run(intro)
        set_run(r, 9.5, MUTED)


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build():
    for directory in (DOCX_PATH.parent, OUT / "pdf", SCREEN, CROPS):
        directory.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((SCREEN / "capture-manifest.json").read_text())
    if manifest.get("databaseConnected") is not False:
        raise RuntimeError("Synthetic capture manifest does not prove databaseConnected=false")
    required = [
        "01-public-home.png", "02-public-search.png", "03-login.png", "04-guest-document.png",
        "05-reader-document.png", "06-reader-saved.png", "07-reader-annotations.png", "08-publisher-news.png",
        "09-publisher-upload.png", "10-admin-dashboard.png", "11-admin-documents.png", "12-admin-reports.png",
    ]
    missing = [name for name in required if not (SCREEN / name).exists()]
    if missing:
        raise FileNotFoundError(f"Missing fixture screenshots: {missing}")

    doc = Document()
    configure(doc)

    # 1. Cover
    cover = doc.add_table(rows=1, cols=1)
    set_table_no_borders(cover)
    cell = cover.cell(0, 0)
    set_cell_shading(cell, "F4F8F4")
    set_cell_margins(cell, top=180, start=220, bottom=180, end=220)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=5)
    r = p.add_run("PAULINIAN ELECTRONIC ARCHIVING SYSTEM")
    set_run(r, 9, GOLD, True)
    add_image(cell, ASSETS / "peas.png", 1.55, "PeAS seal of St. Paul University Dumaguete", None)
    add_text(cell, "User Manual", size=28, color=DEEP_GREEN, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_text(cell, "A concise field guide for finding, reading, publishing, and governing repository content.", size=10.4, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    logo = ASSETS / "office-20of-20research-20-26-20publications-1-.png"
    add_image(cell, logo, 2.65, "Office of Research & Publications wordmark", None)
    add_text(cell, "Office of Research & Publications", size=10, color=GREEN, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    add_text(cell, "St. Paul University Dumaguete", size=9, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=15)
    add_callout(cell, "Edition", "Version 1.0  •  August 2026  •  A5 brochure booklet", fill=PALE_GOLD, size=8.8)
    add_text(cell, "Fictional demonstration data is used throughout the screengrabs.", size=8.3, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    page_break(doc)

    # 2. Overview
    start_page(doc, "01  /  ORIENTATION", "Meet PeAS", "PeAS gives the university community one place to discover, read, and steward research outputs. Access depends on role and on the record's visibility rules.")
    add_two_col(doc, (2.35, 2.65))
    left, right = doc.tables[-1].rows[0].cells
    add_heading(left, "Four audiences", 2)
    add_bullets(left, ["Public visitors browse approved, public records without an account.", "Registered readers save works, use reader tools, and request access where required.", "Content publishers submit news and research files for review.", "Administrators manage governance, permissions, inquiries, users, and reports."], size=8.6)
    add_heading(right, "Access boundary", 2)
    add_bullets(right, ["Public pages show metadata and previews for records marked public.", "Protected full text requires a signed-in reader and an approved request.", "Publisher and administrator screens are role-gated.", "Never infer a permission from a screenshot; confirm the current role and record status."], size=8.6)
    add_callout(doc, "Working model", "Discover → verify metadata → request or open → read responsibly → preserve a clear audit trail.", fill=PALE_GREEN)
    add_image(doc, SCREEN / "01-public-home.png", 5.05, "Fictional PeAS public home screen with search controls", "Figure 1. Public landing page - fictional demonstration data only.", crop=(0, 0, 2880, 1300))
    page_break(doc)

    # 3. Getting started
    start_page(doc, "02  /  FIRST RUN", "Get oriented in three minutes", "Start at the public landing page. The same navigation pattern carries through the reader and administration experiences.")
    add_two_col(doc, (2.6, 2.4))
    left, right = doc.tables[-1].rows[0].cells
    add_heading(left, "Start here", 2)
    add_steps(left, ["Choose Home, News, Contact, or Search from the top navigation.", "Use Search to find a title, author, topic, or keyword.", "Open a record to inspect its abstract, classification, and availability.", "Sign in when you need saved items, reader tools, or a request workflow."], size=8.6)
    add_heading(right, "Account recovery", 2)
    add_steps(right, ["Select Log in, then Forgot password?.", "Enter the account School ID or email address shown on the form.", "Follow the reset link delivered by the institution's configured mail service.", "If the message does not arrive, use the Contact page - do not share a password."], size=8.6)
    add_image(doc, SCREEN / "03-login.png", 4.75, "Fictional PeAS login page with School ID and password fields", "Figure 2. Login and password recovery entry point.", crop=(0, 0, 2880, 1600))
    add_callout(doc, "Safe sign-in", "Use your own account, check the address bar, and sign out on shared computers.", fill=PALE_GOLD, size=8.7)
    page_break(doc)

    # 4. Search
    start_page(doc, "03  /  DISCOVERY", "Search with intent", "The search page is designed for quick narrowing: start broad, then use filters to reduce the result set before opening a record.")
    add_image(doc, SCREEN / "02-public-search.png", 5.1, "Fictional PeAS search hero with repository filters", "Figure 3. Search and filter controls - fictional labels and counts.")
    add_two_col(doc, (2.55, 2.45))
    left, right = doc.tables[-1].rows[0].cells
    add_heading(left, "A reliable sequence", 2)
    add_steps(left, ["Enter a distinctive phrase, author, or topic.", "Choose a document type or classification filter.", "Review the result count and clear filters when the list is too narrow.", "Open the record and confirm the title, author, date, and abstract before citing it."], size=8.6)
    add_heading(right, "Filter cues", 2)
    add_bullets(right, ["Use document type for Thesis, Dissertation, or Confluence.", "Use topic and keyword filters for subject vocabulary.", "Use pagination controls when the result set spans pages.", "A zero-result view is a signal to broaden terms, not a system failure."], size=8.6)
    add_callout(doc, "Caption the evidence", "When sharing a search result, include the date, visible filters, and record status so another reader can reproduce the view.", fill=PALE_BLUE, size=8.6)
    page_break(doc)

    # 5. Record view and request
    start_page(doc, "04  /  RECORDS", "Inspect before you request", "A record page separates descriptive metadata from the protected full text. Read the abstract and availability message before choosing an action.")
    add_image(doc, SCREEN / "04-guest-document.png", 3.28, "Fictional access request dialog for a sample repository record", "Figure 4. Guest access request dialog using fictional contact fields.", crop=(0, 120, 1440, 1450))
    add_two_col(doc, (2.55, 2.45))
    left, right = doc.tables[-1].rows[0].cells
    add_heading(left, "View the record", 2)
    add_steps(left, ["Check the title, author list, publication date, and classification.", "Read the abstract and scan related topics or keywords.", "Use Preview when a public page image is available.", "Choose Request access when the full text is protected."], size=8.5)
    add_heading(right, "Request access", 2)
    add_steps(right, ["Enter a reachable institutional identity in the form.", "Explain the purpose briefly and accurately.", "Submit once; monitor the account or notification route for a decision.", "Use the Contact page if the request remains unresolved."], size=8.5)
    add_callout(doc, "Privacy", "The manual's request form is intentionally empty. Never include passwords, unnecessary personal information, or confidential research details.", fill=PALE_GOLD, size=8.5)
    page_break(doc)

    # 6. Reader PDF controls
    start_page(doc, "05  /  REGISTERED READER", "Read, save, and mark progress", "After sign-in and approval, the reader view adds a focused toolbar for navigation, download, and private study actions.")
    add_image(doc, SCREEN / "05-reader-document.png", 5.08, "Fictional registered-reader PDF viewer showing page controls and reader toolbar", "Figure 5. Reader controls over a fictional sample PDF.", crop=(0, 220, 2272, 1850))
    add_two_col(doc, (2.5, 2.5))
    left, right = doc.tables[-1].rows[0].cells
    add_heading(left, "Toolbar basics", 2)
    add_bullets(left, ["Previous / next page and page number jump.", "Zoom minus, current zoom, plus, Scroll, Full screen.", "Bookmark, Highlight, Area, Note, and Annotations.", "Download PDF exports the approved file when your permission allows it."], size=8.4)
    add_heading(right, "Study loop", 2)
    add_steps(right, ["Save the record to your library.", "Use Mark as read when you finish a pass.", "Add a private note or highlight with enough context to find it again.", "Sign out when finished, especially on shared devices."], size=8.4)
    page_break(doc)

    # 7. Account workspace
    start_page(doc, "06  /  YOUR WORKSPACE", "Keep your research trail tidy", "Saved Items, History, Annotations, and Profile make the registered-reader experience repeatable without changing the source record.")
    cells = add_two_col(doc, (2.45, 2.55))
    left, right = cells
    add_image(left, SCREEN / "06-reader-saved.png", 2.25, "Fictional Saved Items page for Demo Reader", "Saved Items", crop=(0, 0, 2272, 1350))
    add_image(right, SCREEN / "07-reader-annotations.png", 2.25, "Fictional Annotations page with a private demonstration note", "Annotations", crop=(0, 0, 2272, 1500))
    add_two_col(doc, (2.5, 2.5))
    left, right = doc.tables[-1].rows[0].cells
    add_heading(left, "Workspace habits", 2)
    add_bullets(left, ["Saved Items: reopen or remove a record.", "History: revisit recent views and downloads.", "Annotations: filter private notes and highlights.", "Profile: review identity, email verification, and role."], size=8.5)
    add_heading(right, "Account security", 2)
    add_bullets(right, ["Use a unique password and keep recovery details current.", "Review active sessions when the account screen offers them.", "Do not copy a private note into a public inquiry.", "Contact PeAS support from the Contact page for account help."], size=8.5)
    add_callout(doc, "Ownership", "Saved items and annotations are study aids. They do not change the canonical repository record or its classification.", fill=PALE_BLUE, size=8.6)
    page_break(doc)

    # 8. Publisher
    start_page(doc, "07  /  CONTENT PUBLISHER", "Publish with a reviewable trail", "Publishers add news and submit research files. The five-step upload flow collects enough context for a reviewer to validate the record.")
    cells = add_two_col(doc, (2.45, 2.55))
    left, right = cells
    add_image(left, SCREEN / "08-publisher-news.png", 2.25, "Fictional publisher Department News page", "Department News", crop=(0, 0, 2320, 800))
    add_image(right, SCREEN / "09-publisher-upload.png", 2.25, "Fictional publisher upload document wizard at step one", "Upload wizard - step 1 of 5", crop=(0, 0, 2320, 2100))
    add_heading(doc, "Five-step upload workflow", 2)
    add_steps(doc, ["Basics - title, document type, date, and abstract.", "Authors - select existing author records or request a new reference entry.", "Classification - choose research agenda, topics, and keywords.", "Files - attach the approved PDF and supporting assets.", "Review & submit - check the summary, then send to the administrator queue."], size=8.55)
    add_callout(doc, "Before submit", "Use fictional or approved content in demonstrations. A publisher submission is not public until an administrator completes review.", fill=PALE_GOLD, size=8.6)
    page_break(doc)

    # 9. Admin dashboard
    start_page(doc, "08  /  ADMINISTRATOR", "Make the queue actionable", "The administrator dashboard turns pending work into a visible review queue. Metrics are signals for triage, not substitutes for a record-by-record decision.")
    add_image(doc, SCREEN / "10-admin-dashboard.png", 4.95, "Fictional administrator dashboard with demonstration metrics and review cards", "Figure 6. Administrator dashboard - all metrics are synthetic.", crop=(0, 0, 2880, 2050))
    add_two_col(doc, (2.5, 2.5))
    left, right = doc.tables[-1].rows[0].cells
    add_heading(left, "Triage", 2)
    add_steps(left, ["Scan pending uploads and access requests.", "Open the oldest or highest-impact item first.", "Validate metadata, file integrity, and classification.", "Record an explicit decision and reason."], size=8.55)
    add_heading(right, "Dashboard signals", 2)
    add_bullets(right, ["Catalog entries and stored documents describe inventory.", "Page views and visits describe activity coverage.", "Author records help find incomplete references.", "Trending topics are directional and time-bound."], size=8.55)
    add_callout(doc, "Review discipline", "Approve only what you have checked. If a field is uncertain, return the item for clarification rather than guessing.", fill=PALE_GREEN, size=8.6)
    page_break(doc)

    # 10. Governance
    start_page(doc, "09  /  ADMINISTRATION", "Govern the record, not just the file", "PeAS administration is a connected set of controls: metadata quality, permissions, communications, and auditability.")
    add_image(doc, SCREEN / "11-admin-documents.png", 4.9, "Fictional administrator document list with sample statuses and actions", "Figure 7. Document queue and status actions - synthetic records.", crop=(0, 0, 2880, 1650))
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_table_geometry(table, (2.44, 2.44))
    set_repeat_table_header(table.rows[0])
    headers = [("Control area", "Operator focus"), ("Governance check", "Operator focus")]
    for c, (a, b) in zip(table.rows[0].cells, headers):
        set_cell_shading(c, PALE_GREEN)
        set_cell_margins(c, 70, 100, 70, 100)
        p = c.paragraphs[0]
        set_para(p, after=0, line=1.0)
        r = p.add_run(a + "\n")
        set_run(r, 8.3, GREEN, True)
        r = p.add_run(b)
        set_run(r, 8.0, MUTED)
    rows = [
        ("Classification", "Keep agendas, topics, and keywords consistent."),
        ("Authors", "Prefer canonical author records over free-text duplicates."),
        ("Permissions", "Match visibility and download rights to policy."),
        ("Inquiries", "Resolve contact requests without exposing private data."),
        ("Roles / logs", "Grant least privilege; review the audit trail."),
        ("Settings", "Change configuration deliberately and document why."),
    ]
    for left_text, right_text in rows:
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for c, width in zip(cells, (2.44, 2.44)):
            set_cell_width(c, width)
            set_cell_margins(c, 60, 100, 60, 100)
        p = cells[0].paragraphs[0]
        set_para(p, after=0, line=1.0)
        r = p.add_run(left_text)
        set_run(r, 8.1, INK, True)
        p = cells[1].paragraphs[0]
        set_para(p, after=0, line=1.0)
        r = p.add_run(right_text)
        set_run(r, 8.1, INK)
    page_break(doc)

    # 11. Reports and safe use
    start_page(doc, "10  /  OPERATIONS", "Read reports as evidence", "Operational reports help administrators explain repository health, workflow load, and activity coverage. Export only what the audience needs.")
    add_image(doc, SCREEN / "12-admin-reports.png", 4.65, "Fictional operational reports page with synthetic activity metrics", "Figure 8. Reports and exports - synthetic metrics and rankings.", crop=(0, 0, 2880, 2050))
    add_two_col(doc, (2.5, 2.5))
    left, right = doc.tables[-1].rows[0].cells
    add_heading(left, "Common issues", 2)
    add_bullets(left, ["No results: clear filters and broaden the phrase.", "No preview: check visibility and file processing status.", "No download: verify reader approval and permission.", "Stale metric: confirm the report range and coverage note."], size=8.45)
    add_heading(right, "Safe-use reminders", 2)
    add_bullets(right, ["Treat exports as controlled working files.", "Avoid sending personal data in screenshots or reports.", "Cite the record and date when sharing findings.", "Use Contact for support; do not invent a support address."], size=8.45)
    add_callout(doc, "Interpretation", "Every number in this booklet is a fixture example. In production, check the report range, definitions, and last-updated context before acting.", fill=PALE_BLUE, size=8.55)
    page_break(doc)

    # 12. Quick reference
    start_page(doc, "11  /  QUICK REFERENCE", "One-page handoff", "Use this matrix when orienting a new colleague. Role boundaries keep discovery open while protecting review and account actions.")
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_table_geometry(table, (0.95, 1.5, 1.42, 1.01))
    set_repeat_table_header(table.rows[0])
    headings = ["Role", "Can do", "Needs care", "Start here"]
    for cell, text in zip(table.rows[0].cells, headings):
        set_cell_shading(cell, PALE_GREEN)
        set_cell_margins(cell, 70, 80, 70, 80)
        p = cell.paragraphs[0]
        set_para(p, after=0, line=1.0)
        r = p.add_run(text)
        set_run(r, 8.2, GREEN, True)
    matrix = [
        ("Visitor", "Search, filter, view public metadata, request access", "Do not assume protected full text is public", "Home → Search"),
        ("Reader", "Save, annotate, download approved files, mark read", "Keep account and private notes secure", "Log in → Saved Items"),
        ("Publisher", "Create news and submit research files", "Check metadata before review", "Admin → Upload"),
        ("Administrator", "Review, govern, report, manage roles and logs", "Use least privilege and an explicit reason", "Dashboard → Queue"),
    ]
    for row in matrix:
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for c, text, width in zip(cells, row, (0.95, 1.5, 1.42, 1.01)):
            set_cell_width(c, width)
            set_cell_margins(c, 90, 80, 90, 80)
            p = c.paragraphs[0]
            set_para(p, after=0, line=1.0)
            r = p.add_run(text)
            set_run(r, 7.8, INK, bold=(c is cells[0]))
    add_callout(doc, "Support route", "Use the PeAS Contact page for access questions, account recovery, or a suspected metadata issue. Include the record title, action, and time - never a password.", fill=PALE_GOLD, size=8.7)
    add_callout(doc, "Demonstration-data disclaimer", "This booklet uses fixture-served screengrabs only. Names, titles, dates, metrics, classifications, and email addresses are fictional; no production database was connected or read during capture.", fill=PALE_GREEN, size=8.8)
    add_text(doc, "PeAS User Manual  •  Version 1.0  •  August 2026  •  Current implementation guide", size=8.4, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
